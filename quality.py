"""Pure, testable validation and explicitly selected transformations."""
import csv
import io
import re
from datetime import datetime, date
import pandas as pd
from openpyxl import Workbook, load_workbook

MAX_ROWS=50000
MAX_COLUMNS=150
EMAIL=re.compile(r'^[^\s@]+@[^\s@.]+(?:\.[^\s@.]+)+$')

def missing(value):
    return pd.isna(value) or str(value).strip()==''

def as_text(value):
    if pd.isna(value): return ''
    if isinstance(value,(date,datetime)): return value.strftime('%Y-%m-%d')
    return str(value)

def sheets(content):
    book=load_workbook(io.BytesIO(content),read_only=True)
    try: return book.sheetnames
    finally: book.close()

def load_table(content,name,sheet=0,delimiter=','):
    if len(content)>10*1024*1024: raise ValueError('Dosya en fazla 10 MB olabilir.')
    if name.lower().endswith('.csv'):
        text=content.decode('utf-8-sig')
        rows=list(csv.reader(io.StringIO(text),delimiter=delimiter))
        if not rows: raise ValueError('Dosya boş.')
        headers=rows[0]
        if any(len(row)!=len(headers) for row in rows[1:]): raise ValueError('Satırlarda sütun sayısı tutarsız. Ayırıcıyı kontrol et.')
        frame=pd.DataFrame(rows[1:],columns=headers)
    elif name.lower().endswith('.xlsx'):
        # Explicitly values-only: the export is a new flat table, not a workbook editor.
        raw=pd.read_excel(io.BytesIO(content),sheet_name=sheet,header=None,dtype=object,keep_default_na=False,nrows=MAX_ROWS+2)
        if raw.empty: raise ValueError('Seçilen sayfa boş.')
        headers=[as_text(x) for x in raw.iloc[0]]
        frame=raw.iloc[1:].copy();frame.columns=headers;frame=frame.reset_index(drop=True).map(as_text)
    else: raise ValueError('Yalnızca CSV ve XLSX destekleniyor.')
    if any(not h.strip() for h in frame.columns) or len(set(frame.columns))!=len(frame.columns):
        raise ValueError('Sütun başlıkları dolu ve benzersiz olmalı.')
    if len(frame)>MAX_ROWS or len(frame.columns)>MAX_COLUMNS: raise ValueError('Sınır: 50.000 satır ve 150 sütun.')
    return frame

def parse_date(value,format):
    if isinstance(value,(date,datetime)): return value.strftime('%Y-%m-%d')
    text=str(value).strip()
    parsed=datetime.strptime(text,format)
    if parsed.strftime(format)!=text: raise ValueError('Tarih biçimi tam eşleşmiyor.')
    return parsed.strftime('%Y-%m-%d')

def inspect(frame,email_columns=(),date_columns=(),date_format='%Y-%m-%d'):
    issues=[]
    def add(row,column,kind,value):issues.append({'Satır':row+2,'Sütun':column,'Sorun':kind,'Değer':str(value)})
    for pos,(_,row) in enumerate(frame.iterrows()):
        for col,value in row.items():
            if missing(value):add(pos,col,'Eksik hücre',value);continue
            if col in email_columns and not EMAIL.fullmatch(str(value)):add(pos,col,'E-posta biçimi',value)
            if col in date_columns:
                try:parse_date(value,date_format)
                except (ValueError,TypeError):add(pos,col,'Tarih biçimi',value)
    for pos,duplicate in enumerate(frame.duplicated(keep='first')):
        if duplicate:add(pos,'(tüm satır)','Yinelenen satır','')
    return pd.DataFrame(issues,columns=['Satır','Sütun','Sorun','Değer'])

def transform(frame,*,trim=False,drop_duplicates=False,fill_columns=(),fill_value='',date_columns=(),date_format='%Y-%m-%d',normalize_dates=False):
    result=frame.copy(deep=True)
    if trim:result=result.map(lambda v:v.strip() if isinstance(v,str) else v)
    for col in fill_columns:result[col]=result[col].map(lambda v:fill_value if missing(v) else v)
    if normalize_dates:
        def normalize(v):
            try:return parse_date(v,date_format)
            except (ValueError,TypeError):return v
        for col in date_columns:result[col]=result[col].map(normalize)
    if drop_duplicates:result=result.drop_duplicates(keep='first')
    return result.reset_index(drop=True)

def csv_bytes(frame):
    def safe(value):
        text=as_text(value)
        return "'"+text if text.lstrip().startswith(('=','+','-','@')) or text.startswith(('\t','\r','\n')) else text
    result=frame.map(safe)
    result.columns=[safe(c) for c in frame.columns]
    return result.to_csv(index=False).encode('utf-8-sig')

def xlsx_bytes(frame):
    book=Workbook();sheet=book.active;sheet.title='Clean data'
    for row in [list(frame.columns),*frame.values.tolist()]:
        sheet.append([as_text(value) for value in row])
        for cell in sheet[sheet.max_row]: cell.data_type='s'
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    output=io.BytesIO();book.save(output);return output.getvalue()
